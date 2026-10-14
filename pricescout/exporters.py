"""Export scraped products to CSV, JSON, formatted Excel, Google Sheets and MongoDB."""

from __future__ import annotations

import json
import math
from pathlib import Path
from typing import Iterable

import pandas as pd

from .models import EXPORT_COLUMNS, Product

SUPPORTED_FORMATS = ("csv", "json", "xlsx")


def products_to_df(products: Iterable[Product]) -> pd.DataFrame:
    rows = []
    for p in products:
        row = p.model_dump()
        row["scraped_at"] = p.scraped_at.strftime("%Y-%m-%d %H:%M:%S UTC")
        rows.append(row)
    df = pd.DataFrame(rows, columns=EXPORT_COLUMNS)
    if not df.empty:
        df["in_stock"] = df["in_stock"].astype(bool)
    return df


def category_summary(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame()
    data = df.assign(category=df["category"].fillna("Unknown"))
    summary = (
        data.groupby("category")
        .agg(
            products=("title", "count"),
            avg_price=("price", "mean"),
            min_price=("price", "min"),
            max_price=("price", "max"),
            avg_rating=("rating", "mean"),
            in_stock_pct=("in_stock", "mean"),
            units_in_stock=("stock_count", "sum"),
        )
        .reset_index()
        .sort_values("products", ascending=False)
    )
    summary["in_stock_pct"] = (summary["in_stock_pct"] * 100).round(1)
    return summary.round({"avg_price": 2, "min_price": 2, "max_price": 2, "avg_rating": 2})


# --------------------------------------------------------------------- files
def export_csv(df: pd.DataFrame, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False, encoding="utf-8-sig")  # BOM -> Excel opens UTF-8 correctly
    return path


def export_json(df: pd.DataFrame, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    records = json.loads(df.to_json(orient="records", force_ascii=False))
    path.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def export_excel(df: pd.DataFrame, path, changes: pd.DataFrame | None = None):
    """Client-ready workbook: Products, Summary by category and (optional) Changes.

    ``path`` can be a file path or a binary buffer (``io.BytesIO``) for downloads.
    """
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if isinstance(path, (str, Path)):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
    sheets = {"Products": df, "Summary": category_summary(df)}
    if changes is not None and not changes.empty:
        sheets["Changes"] = changes

    header_fill = PatternFill("solid", fgColor="0F4C5C")
    header_font = Font(color="FFFFFF", bold=True)
    money_cols = {"price", "avg_price", "min_price", "max_price", "old_price", "new_price", "price_diff"}

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            # Print-ready: landscape, all columns on one page width
            ws.page_setup.orientation = "landscape"
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
            if frame.shape[1]:
                ws.auto_filter.ref = ws.dimensions
            for idx, column in enumerate(frame.columns, start=1):
                cell = ws.cell(row=1, column=idx)
                cell.fill, cell.font = header_fill, header_font
                cell.alignment = Alignment(vertical="center")
                values = ["" if pd.isna(v) else str(v) for v in frame[column].head(200).tolist()]
                width = min(60, max([len(str(column))] + [len(v) for v in values]) + 2)
                ws.column_dimensions[get_column_letter(idx)].width = width
                if column in money_cols:
                    for row in range(2, len(frame) + 2):
                        ws.cell(row=row, column=idx).number_format = "#,##0.00"
    return path


def export_all(df: pd.DataFrame, out_dir: str | Path, formats: Iterable[str],
               basename: str = "products", changes: pd.DataFrame | None = None) -> dict[str, Path]:
    out_dir = Path(out_dir)
    written: dict[str, Path] = {}
    for fmt in formats:
        fmt = fmt.strip().lower()
        if fmt == "csv":
            written["csv"] = export_csv(df, out_dir / f"{basename}.csv")
        elif fmt == "json":
            written["json"] = export_json(df, out_dir / f"{basename}.json")
        elif fmt in {"xlsx", "excel"}:
            written["xlsx"] = export_excel(df, out_dir / f"{basename}.xlsx", changes)
        elif fmt:
            raise ValueError(f"Unsupported format '{fmt}'. Choose from {SUPPORTED_FORMATS}.")
    return written


# ------------------------------------------------------------ integrations
def _sheet_value(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    if hasattr(value, "item"):  # numpy scalar -> python
        return value.item()
    return value


def export_google_sheets(df: pd.DataFrame, spreadsheet_key: str, worksheet: str = "Products",
                         credentials_file: str | None = None, client=None) -> int:
    """Replace a worksheet with the scraped data. Returns the number of rows written.

    Uses a Google service account (share the sheet with the service account email).
    """
    if client is None:
        import gspread

        client = gspread.service_account(filename=credentials_file) if credentials_file else gspread.service_account()
    book = client.open_by_key(spreadsheet_key)
    try:
        ws = book.worksheet(worksheet)
    except Exception:  # gspread.WorksheetNotFound
        ws = book.add_worksheet(title=worksheet, rows=max(100, len(df) + 10), cols=max(10, df.shape[1] + 2))
    values = [list(df.columns)] + [[_sheet_value(v) for v in row] for row in df.itertuples(index=False)]
    ws.clear()
    ws.update(range_name="A1", values=values)
    return len(values) - 1


def export_mongodb(df: pd.DataFrame, uri: str = "mongodb://localhost:27017", database: str = "pricescout",
                   collection: str = "products", client=None) -> int:
    """Upsert products by key (UPC, else URL). Returns the number of documents written."""
    if client is None:
        from pymongo import MongoClient

        client = MongoClient(uri, serverSelectionTimeoutMS=5000)
    coll = client[database][collection]
    written = 0
    for record in json.loads(df.to_json(orient="records", force_ascii=False)):
        key = record.get("upc") or record["url"]
        # The product key is the document _id, so re-running a scrape updates instead of duplicating.
        coll.replace_one({"_id": key}, {"_id": key, **record}, upsert=True)
        written += 1
    return written
